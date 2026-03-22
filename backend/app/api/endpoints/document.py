import os
import io
import asyncio
from fastapi import APIRouter, HTTPException, UploadFile, File, Response
from fastapi.responses import StreamingResponse
from typing import List
from app.schemas.document import DocumentResponse, DocumentUpdate
from app.db.supabase import supabase
from app.services.ocr_service import ocr_service
from app.core.config import settings

router = APIRouter()



ALLOWED_MIME_TYPES = [
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
    "application/pdf",
    "image/png",
    "image/jpeg"
]

# Types that need PaddleOCR
OCR_TYPES = ["application/pdf", "image/png", "image/jpeg"]
# Types that can be read directly
DIRECT_TYPES = ["text/plain", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"]


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text from a plain text file."""
    # Try common encodings
    for encoding in ['utf-8', 'gbk', 'gb2312', 'gb18030', 'latin-1']:
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return file_bytes.decode('utf-8', errors='replace')


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a Word .docx file using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        
        # Extract tables as well
        for table in doc.tables:
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_data.append(cell.text.strip().replace('\n', ' '))
                if row_data:
                    paragraphs.append(" | ".join(row_data))
                    
        return "\n\n".join(paragraphs)
    except ImportError:
        # Fallback: extract from XML directly
        import zipfile
        import re
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                xml_content = z.read('word/document.xml').decode('utf-8')
                # Simple regex to extract text between XML tags
                text = re.sub(r'<[^>]+>', '', xml_content)
                # Clean up whitespace
                lines = [line.strip() for line in text.split('\n') if line.strip()]
                return "\n\n".join(lines)
        except Exception as e:
            return f"[无法解析Word文档: {str(e)}]"
    except Exception as e:
        return f"[Word文档解析错误: {str(e)}]"


@router.post("/upload", response_model=DocumentResponse)
async def upload_document(file: UploadFile = File(...)):
    if file.content_type not in ALLOWED_MIME_TYPES:
        raise HTTPException(status_code=400, detail="不支持的文件格式")
    
    file_bytes = await file.read()
    file_size = len(file_bytes)
    if file_size > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件大小超过限制 (10MB)")

    # 1. Insert record into Supabase with status=pending
    doc_data = {
        "file_name": file.filename,
        "file_type": file.content_type,
        "file_size": file_size,
        "status": "pending"
    }
    db_response = supabase.table("documents").insert(doc_data).execute()
    if not db_response.data:
        raise HTTPException(status_code=500, detail="数据库插入失败")
    
    doc_id = db_response.data[0]["id"]
    
    # 2. Save the file to Supabase Storage
    ext = os.path.splitext(file.filename)[1]
    if not ext:
        if "pdf" in file.content_type: ext = ".pdf"
        elif "png" in file.content_type: ext = ".png"
        elif "jpeg" in file.content_type or "jpg" in file.content_type: ext = ".jpg"
        elif "word" in file.content_type: ext = ".docx"
        else: ext = ".txt"
        
    storage_path = f"{doc_id}{ext}"
    try:
        supabase.storage.from_(settings.SUPABASE_STORAGE_BUCKET).upload(
            storage_path, 
            file_bytes, 
            {"content-type": file.content_type}
        )
    except Exception as e:
        print(f"Failed to save file to Supabase Storage: {e}")
        # Rollback db insert if possible or mark as failed
        supabase.table("documents").delete().eq("id", doc_id).execute()
        raise HTTPException(status_code=500, detail=f"文件存储失败: {str(e)}")

    # 3. Process based on file type
    if file.content_type in DIRECT_TYPES:
        # TXT / DOCX: extract content directly, no OCR needed
        try:
            if file.content_type == "text/plain":
                content = extract_text_from_txt(file_bytes)
            else:
                content = extract_text_from_docx(file_bytes)
            
            update_res = supabase.table("documents").update({
                "status": "done",
                "ocr_content": content
            }).eq("id", doc_id).execute()
            return update_res.data[0]
        except Exception as e:
            supabase.table("documents").update({
                "status": "failed",
                "ocr_content": f"内容提取失败: {str(e)}"
            }).eq("id", doc_id).execute()
            raise HTTPException(status_code=500, detail=str(e))
    else:
        # PDF / Image: just save, user triggers OCR manually
        update_res = supabase.table("documents").update({
            "status": "uploaded"
        }).eq("id", doc_id).execute()
        return update_res.data[0]


@router.post("/ocr/{doc_id}", response_model=DocumentResponse)
async def trigger_ocr(doc_id: str):
    """Manually trigger OCR for a document (PDF/Image only)."""
    # 1. Get document info
    doc_res = supabase.table("documents").select("*").eq("id", doc_id).execute()
    if not doc_res.data:
        raise HTTPException(status_code=404, detail="文档未找到")
    
    doc = doc_res.data[0]
    file_type = doc["file_type"]
    
    if file_type not in OCR_TYPES:
        raise HTTPException(status_code=400, detail="该文件类型不需要OCR")
    
    # 2. Read the file from Supabase Storage
    ext = os.path.splitext(doc["file_name"])[1]
    if not ext:
        if "pdf" in file_type: ext = ".pdf"
        elif "png" in file_type: ext = ".png"
        elif "jpeg" in file_type or "jpg" in file_type: ext = ".jpg"
    
    storage_path = f"{doc_id}{ext}"
    try:
        file_bytes = supabase.storage.from_(settings.SUPABASE_STORAGE_BUCKET).download(storage_path)
    except Exception as e:
        raise HTTPException(status_code=404, detail=f"存储库中未找到文件: {str(e)}")
    
    # 3. Update status to running
    supabase.table("documents").update({"status": "running"}).eq("id", doc_id).execute()
    
    # 4. Submit to PaddleOCR
    try:
        job_id = await ocr_service.submit_job(file_bytes, doc["file_name"])
    except Exception as e:
        supabase.table("documents").update({
            "status": "failed",
            "ocr_content": f"提交OCR任务失败: {str(e)}"
        }).eq("id", doc_id).execute()
        raise HTTPException(status_code=500, detail=str(e))
    
    update_res = supabase.table("documents").update({
        "status": "running",
        "job_id": job_id
    }).eq("id", doc_id).execute()
    return update_res.data[0]


@router.get("/list", response_model=List[DocumentResponse])
def list_documents():
    response = supabase.table("documents").select("*").order("created_at", desc=True).execute()
    return response.data


@router.put("/{doc_id}", response_model=DocumentResponse)
def update_document_content(doc_id: str, doc_in: DocumentUpdate):
    update_data = {
        "ocr_content": doc_in.ocr_content
    }
    response = supabase.table("documents").update(update_data).eq("id", doc_id).execute()
    if not response.data:
        raise HTTPException(status_code=404, detail="文档未找到")
    return response.data[0]


@router.delete("/{doc_id}")
def delete_document(doc_id: str):
    # 1. Delete associated chats first (foreign key)
    try:
        supabase.table("chats").delete().eq("doc_id", doc_id).execute()
    except Exception as e:
        print(f"Warning: chats cleanup for {doc_id}: {e}")
    
    # 2. Delete the document
    try:
        response = supabase.table("documents").delete().eq("id", doc_id).execute()
        print(f"Delete response for {doc_id}: data={response.data}")
    except Exception as e:
        print(f"Error deleting document {doc_id}: {e}")
        raise HTTPException(status_code=500, detail=f"删除失败: {str(e)}")
    
    # 3. Clean up storage file
    for ext in ['.pdf', '.png', '.jpg', '.jpeg', '.docx', '.txt']:
        storage_path = f"{doc_id}{ext}"
        try:
            # remove() takes a list of paths
            supabase.storage.from_(settings.SUPABASE_STORAGE_BUCKET).remove([storage_path])
        except Exception:
            pass
                
    return {"message": "文档删除成功"}
    

@router.get("/file/{doc_id}")
async def get_document_file(doc_id: str):
    """Fetch the document file from Supabase Storage."""
    # 1. Get document info to find full filename/type
    doc_res = supabase.table("documents").select("*").eq("id", doc_id).execute()
    if not doc_res.data:
        raise HTTPException(status_code=404, detail="文档未找到")
    
    doc = doc_res.data[0]
    file_type = doc["file_type"]
    ext = os.path.splitext(doc["file_name"])[1]
    if not ext:
        if "pdf" in file_type: ext = ".pdf"
        elif "png" in file_type: ext = ".png"
        elif "jpeg" in file_type or "jpg" in file_type: ext = ".jpg"
        elif "word" in file_type: ext = ".docx"
        else: ext = ".txt"
        
    storage_path = f"{doc_id}{ext}"
    try:
        file_bytes = supabase.storage.from_(settings.SUPABASE_STORAGE_BUCKET).download(storage_path)
        return StreamingResponse(io.BytesIO(file_bytes), media_type=file_type)
    except Exception as e:
        raise HTTPException(status_code=404, detail=f"文件不存在: {str(e)}")
